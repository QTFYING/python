import openpyxl
import pandas as pd

def load_workbook(file_path):
    return openpyxl.load_workbook(file_path, data_only=False)

def get_hyperlink(cell):
    return cell.hyperlink.target if cell.hyperlink else None

def extract_movie_links(ws):
    movie_links = {}
    for row in ws.iter_rows(min_row=2, values_only=False):
        movie_cell = row[0]
        baidu_cell = row[1]
        kuake_cell = row[2]
        xunlei_cell = row[3]

        movie_name = movie_cell.value

        baidu_link = get_hyperlink(baidu_cell)
        kuake_link = get_hyperlink(kuake_cell)
        xunlei_link = get_hyperlink(xunlei_cell)

        movie_links[movie_name] = {
            '百度网盘': baidu_link,
            '夸克网盘': kuake_link,
            '迅雷云盘': xunlei_link
        }
    return movie_links

def print_movie_links(movie_links):
    for movie, links in movie_links.items():
        print(f"电影名称: {movie}")
        if links['百度网盘']:
            print(f"百度网盘: {links['百度网盘']}")
        if links['夸克网盘']:
            print(f"夸克网盘: {links['夸克网盘']}")
        if links['迅雷云盘']:
            print(f"迅雷云盘: {links['迅雷云盘']}")
        print("-" * 40)

def main():
    file_path = 'movie_list_less.xlsx'
    wb = load_workbook(file_path)
    ws = wb.active

    movie_links = extract_movie_links(ws)
    print_movie_links(movie_links)

if __name__ == "__main__":
    main()